"""导出 xlsx 的**每一格**都必须与同一窗口的接口 JSON 一致（2026-09-23 第 11 轮）。

## 为什么单独立一条

导出是**给外部看的凭证**（用户拿它对账、存档、发给别人），而它与页面走的是两条完全不同的
序列化路径：页面走 JSON，导出走 `ws.append([...])` 一列一列铺。两边一旦分叉，**没有任何接口会报错**：

| 写坏的方式 | 表现 |
|---|---|
| `ws.append` 里两个标签顺序写反 | 表里「金额」那一列其实是运费 —— 数看着都合理，人对不出来 |
| 文件名按**锚点日**命名、内容按整月取数 | 拿文件名找回来的凭证「名字与内容不符」（2026-09-19 R2-4 那一类） |
| 命中不了成本的行印成数字 `0` | 那一格看起来像「这单没成本、毛利全额」（2026-09-19 第十七轮修掉的老缺陷又回来了） |
| 金额写成文本 | 用户在 Excel 里选一列 `SUM` 得 0（R13-R7），账得拿计算器重算 |

## 这条测试为什么不是空转（反空转判据写在断言里）

- 用**真实链路**造一张当日已送达的单，其中**一行有成本、一行没成本** → 商品页必须出现一格 `—`；
- 同一窗口分别打 `/reports/turnover`、`/reports/products`（JSON）与 `/reports/export`（xlsx），
  **逐格比对**（金额量化到两位小数后必须相等），系列行 / 商品行逐行相等；
- 断言窗口里至少有 1 张单、商品行至少 1 行、且**必须出现过一次 `—`** —— 否则这条判据对
  "算不出成本的行"这件事没有发言权（表里没这种行时它恒真）。

⚠️ 只比"导出 == 接口"，不比绝对值：同一个测试库里还有别的用例留下的单，
绝对值每天都在变，那样的断言要么假红要么被写成恒真。
"""

from __future__ import annotations

import random
from datetime import timedelta
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

import pytest
from sqlalchemy import delete, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from starlette.testclient import TestClient

from app.core.business_time import business_today
from app.models import (
    CashFlow,
    DriverBill,
    InventoryMovement,
    Ledger,
    OperationLog,
    Order,
    OrderProduct,
    Place,
    Product,
)
from tests.conftest import auth_headers

CENTS = Decimal("0.01")


def _uniq(prefix: str) -> str:
    return f"{prefix}-{random.randint(100000, 999999)}"


def _money(v) -> Decimal:
    """把接口给的金额/单元格里的数统一成"两位小数的 Decimal"再比（浮点与 str 都能进来）。

    ⚠️ 必须显式 `ROUND_HALF_UP`：导出那边（`reports._money`）就是 HALF_UP，
    而 `Decimal.quantize` 的**默认是 ROUND_HALF_EVEN** —— 第三位是 5 时两边差一分
    （实测：接口 `12.3450` → 默认舍入给 12.34、导出给 12.35，断言假红）。
    """
    return Decimal(str(v)).quantize(CENTS, rounding=ROUND_HALF_UP)


def _mk_product(client: TestClient, h: dict[str, str], *, cost: str) -> dict:
    r = client.post(
        "/api/v1/products",
        json={
            "name": _uniq("导出对账商品"),
            "default_unit_price": "40",
            "cost_price": cost,
            "unit": "件",
            "stock": 0,
            "category": _uniq("导出对账分类"),
        },
        headers=h,
    )
    assert r.status_code == 201, r.text
    return r.json()


def _mk_driver(client: TestClient, h: dict[str, str], **extra) -> tuple[int, str]:
    """专用司机（不污染 fixture 里那个司机 —— 给他挂规则会改变同一轮其它用例的行为）。

    `**extra` 直接进建号请求（例如 `billing_mode="SALARY", salary="5000"` 造一个工资制司机）。
    """
    import uuid

    phone = "13" + uuid.uuid4().hex[:9].translate(str.maketrans("abcdef", "012345"))
    body = {"phone": phone, "password": "pass12345", "full_name": "导出对账专用司机", "role": "driver"}
    body.update(extra)
    r = client.post("/api/v1/users", json=body, headers=h)
    assert r.status_code in (200, 201), r.text
    driver_id = int(r.json()["id"])
    r = client.post("/api/v1/auth/login", json={"phone": phone, "password": "pass12345"})
    assert r.status_code == 200, r.text
    return driver_id, r.json()["access_token"]


def _span_label(start, end) -> str:
    """导出表头那段窗口标签的**独立重算**（不 import 被测函数，否则断言是自证）。

    规则来自 `reports._span_label` 的文档字符串（那份文档是 2026-09-22 定稿的）：
    整天 `9-18` / 整月 `2026-09月` / 其余一段 `9-1~9-20`。
    ⚠️ 只有"整月"那一支带年份 —— 这条是**当前**行为，钉住它意味着以后要改得是故意的；
    另见台账「待拍板」里那条"导出凭证的区间标签要不要都带年份"。
    """
    if start == end:
        return f"{start.month}-{start.day}"
    if start.day == 1 and (start.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1) == end:
        return f"{start.year}-{start.month:02d}月"
    return f"{start.month}-{start.day}~{end.month}-{end.day}"


def _xlsx(client: TestClient, h: dict[str, str], **params) -> tuple[object, str]:
    """打导出接口 → (工作表字典, 文件名)。文件名从 `Content-Disposition` 里取。"""
    from openpyxl import load_workbook

    r = client.get("/api/v1/reports/export", params=params, headers=h)
    assert r.status_code == 200, r.text
    disposition = r.headers.get("content-disposition", "")
    assert "filename=" in disposition, f"导出没给文件名：{disposition!r}"
    name = disposition.split("filename=", 1)[1].strip().strip('"')
    wb = load_workbook(BytesIO(r.content), data_only=True)
    return {ws.title: ws for ws in wb.worksheets}, name


def _cells(ws, row: int) -> list:
    """取一行里**有值的那些格**（`ws.append([...])` 写出来的行，末尾空格不算）。"""
    return [c.value for c in ws[row] if c.value is not None]


def _rows_after(ws, header_row: int) -> list[list]:
    """取表头下面**到下一个空行/下一段为止**的所有行（不是"按接口的长度去读那么多行"）。

    ⚠️ 这里必须**数表里真实的行**：第一版写成
    `[_cells(ws, header_row + i) for i in range(1, len(items) + 1)]` —— 循环次数由接口决定，
    于是"导出少写了几行"这件事**永远看不出来**（注入"曲线只写第一行"照样全绿，实测抓到）。
    """
    out: list[list] = []
    row = header_row + 1
    while row <= ws.max_row:
        got = _cells(ws, row)
        if not got:                     # 空行 = 这一段结束
            break
        out.append(got)
        row += 1
    return out


def _cols(ws, row: int, n: int) -> list:
    """按**列号**取一行前 n 格（空值/空串都保留）。

    ⚠️ 为什么不能用"有值才取"的 `_cells` 去比固定列的表：某一格是空串时那一行会比表头短，
    按下标取值就 `IndexError`（全量跑时被别的用例留下的"异常原因为空"的单抓到了）。
    固定列的表一律用这个；"末尾列本来就没写"的行才用 `_cells`。
    """
    return [ws.cell(row=row, column=i + 1).value for i in range(n)]


def _rows_cols(ws, header_row: int, n: int) -> list[list]:
    """表头下面每一行按列号取前 n 格，直到**整行都空**为止。"""
    out: list[list] = []
    row = header_row + 1
    while row <= ws.max_row:
        vals = _cols(ws, row, n)
        if all(v is None for v in vals):
            break
        out.append(vals)
        row += 1
    return out


def _cleanup(db: Session, order_ids: list[int], product_ids: list[int]) -> None:
    """把自己造的探针数据清掉（这个库是多个用例共用的，留着会影响别人的绝对值断言）。"""
    if not order_ids and not product_ids:
        return
    if order_ids:
        db.execute(update(Place).where(Place.first_order_id.in_(order_ids)).values(first_order_id=None))
        db.execute(delete(DriverBill).where(DriverBill.order_id.in_(order_ids)))
        # ⚠️ 现金流水也要删：收现金的单会写一条 `cash_flows`，而它挂着 `order_id` 外键 ——
        #    不先删它，下面那句 `DELETE FROM orders` 会 IntegrityError（第 12 轮实测踩到）。
        db.execute(delete(CashFlow).where(CashFlow.order_id.in_(order_ids)))
        db.execute(delete(Ledger).where(Ledger.order_id.in_(order_ids)))
        db.execute(delete(OperationLog).where(OperationLog.order_id.in_(order_ids)))
        db.execute(delete(InventoryMovement).where(InventoryMovement.order_id.in_(order_ids)))
        db.execute(delete(OrderProduct).where(OrderProduct.order_id.in_(order_ids)))
        db.execute(delete(Order).where(Order.id.in_(order_ids)))
    if product_ids:
        db.execute(delete(InventoryMovement).where(InventoryMovement.product_id.in_(product_ids)))
        try:
            db.execute(delete(Product).where(Product.id.in_(product_ids)))
            db.commit()
        except IntegrityError:
            # 商品还被别的表指着（价格规则/可见性/成本时间轴…）→ 退回**软删**：
            # 硬删不是这条测试的目的，为它跟外键较劲不值得（软删后列表与下单都看不到它）。
            db.rollback()
            db.execute(update(Product).where(Product.id.in_(product_ids)).values(is_deleted=True))
            db.commit()
    if order_ids:
        db.commit()


@pytest.mark.dispatcher
@pytest.mark.fast
@pytest.mark.regression
def test_export_cells_match_api_for_the_same_window(
    client: TestClient, token_dispatcher: str, db_session: Session, users: dict
) -> None:
    h = auth_headers(token_dispatcher)
    today = business_today()
    anchor = today.isoformat()

    # ---- 造现场：一行有成本（毛利算得出来）、一行没成本（毛利那一格必须是「—」）----
    with_cost = _mk_product(client, h, cost="6.4")
    no_cost = _mk_product(client, h, cost="0")
    driver_id, dtoken = _mk_driver(client, h)
    r = client.post(
        "/api/v1/orders",
        json={
            "shipper_id": users["shipper"].id,
            "lines": [
                {
                    "product_id": with_cost["id"],
                    "product_name_snapshot": with_cost["name"],
                    "quantity": 3,
                    "unit_price": "40",
                    "line_total": "120",
                },
                {
                    "product_id": no_cost["id"],
                    "product_name_snapshot": no_cost["name"],
                    "quantity": 1,
                    "unit_price": "50",
                    "line_total": "50",
                },
            ],
            "address_detail": "导出逐格对账路 1 号",
            "freight_fee": "30",
        },
        headers=h,
    )
    assert r.status_code in (200, 201), r.text
    order = r.json()
    order_ids = [int(order["id"])]
    product_ids = [int(with_cost["id"]), int(no_cost["id"])]
    try:
        assert client.post(
            f"/api/v1/orders/{order['id']}/assign", json={"driver_id": driver_id}, headers=h
        ).status_code in (200, 201)
        hd = auth_headers(dtoken)
        assert client.post(f"/api/v1/orders/{order['id']}/driver-ack", headers=hd).status_code in (200, 201)
        r = client.post(
            f"/api/v1/orders/{order['id']}/complete",
            json={
                "payment": "arrears",
                # 送达照片按规矩必须给（免拍照只给"不按单拿钱"的工资制司机 —— 那取决于
                # 这个探针司机当时的计费快照，拿它当前提这条测试就会随别处的改动飘）。
                # 这里给一张**形状正确**的 URL：后端刻意只校验"是本系统上传端点的产物"
                # （不校验文件是否存在，`order_flow.complete_delivery` 里写了理由）。
                "delivery_photo_urls": [f"/static/uploads/delivery/{order['id']}/export-probe.jpg"],
            },
            headers=hd,
        )
        assert r.status_code in (200, 201), r.text

        # ---- 同一窗口：接口 JSON ----
        q = {"mode": "day", "date": anchor}
        turnover = client.get("/api/v1/reports/turnover", params=q, headers=h)
        assert turnover.status_code == 200, turnover.text
        t = turnover.json()
        products = client.get("/api/v1/reports/products", params=q, headers=h)
        assert products.status_code == 200, products.text
        p = products.json()

        # 反空转：窗口里真的有单、商品行真的有、且商品行里真的有一行算不出成本
        assert t["total_orders"] >= 1, "刚送达的单没进当天窗口 —— 后面的逐格比对没有意义"
        assert p["items"], "商品经营一行都没有"
        assert any((it.get("covered_lines") or 0) == 0 for it in p["items"]), (
            "这个窗口里没有『算不出成本』的商品行 —— 「—」那一格根本没被覆盖，判据会恒真"
        )

        # ---- 导出（营业纵览）----
        sheets, name = _xlsx(client, h, kind="turnover", **q)
        assert "营业纵览" in sheets, f"导出的工作表叫 {list(sheets)}"
        ws = sheets["营业纵览"]

        # ① 文件名 = **这一份报表真实取数的区间**（mode=day 时就是那一天）
        assert name == f"turnover-report-{anchor}.xlsx", (
            f"文件名与取数区间不一致：{name}；按文件名找回来的凭证会与内容对不上（R2-4）"
        )
        # ② 第 1 行的窗口标签就是这一段（格式见 `_span_label` 的文档）
        head = _cells(ws, 1)
        assert head[0] == "营业纵览", head
        assert head[1] == _span_label(today, today), f"表头写的窗口不是请求的那一段：{head}"
        assert head[2] == "金额口径：已送达未撤销", head

        # ③ 逐格比对（标签 → 接口字段）
        r2 = _cells(ws, 2)
        assert r2[0] == "营业金额" and _money(r2[1]) == _money(t["total_amount"]), (r2, t["total_amount"])
        assert r2[2] == "订单数" and int(r2[3]) == int(t["total_orders"]), (r2, t["total_orders"])
        assert r2[4] == "单均价" and _money(r2[5]) == _money(t["avg_order"]), (r2, t["avg_order"])
        r3 = _cells(ws, 3)
        assert r3[0] == "司机运费支出(按计费规则应付)", r3
        assert _money(r3[1]) == _money(t["total_freight"]), (r3[1], t["total_freight"])
        assert r3[2] == "商品毛利(仅算得出成本的行)", r3
        # 毛利两侧必须同一批行：cost_covered_amount − cost_total
        assert _money(r3[3]) == _money(Decimal(str(t["cost_covered_amount"])) - Decimal(str(t["cost_total"]))), (
            r3[3],
            t["cost_covered_amount"],
            t["cost_total"],
        )
        r4 = _cells(ws, 4)
        assert r4[0] == "货损件数" and int(r4[1]) == int(t["damage_qty"]), (r4, t["damage_qty"])
        assert r4[2] == "货损金额" and _money(r4[3]) == _money(t["damage_amount"]), (r4, t["damage_amount"])
        assert r4[4] == "已收" and _money(r4[5]) == _money(t["collected"]), (r4, t["collected"])
        r5 = _cells(ws, 5)
        assert r5[0] == "挂账未收" and _money(r5[1]) == _money(t["arrears_total"]), (r5, t["arrears_total"])
        assert r5[2] == "已撤销订单" and int(r5[3]) == int(t["cancelled_orders"]), (r5, t["cancelled_orders"])

        # ④ 曲线那一段：表头 + 逐行（时间/单数/金额/运费）——**行数从表里数出来**
        series_head_row = next(row for row in range(1, ws.max_row + 1) if _cells(ws, row)[:1] == ["时间"])
        assert _cells(ws, series_head_row) == ["时间", "单数", "金额", "运费"], _cells(ws, series_head_row)
        got_series = _rows_cols(ws, series_head_row, 4)
        assert len(got_series) == len(t["series"]), (
            f"导出写了 {len(got_series)} 行、接口给了 {len(t['series'])} 行"
        )
        for row, item in zip(got_series, t["series"]):
            assert row[0] == item["label"], (row, item)
            assert int(row[1]) == int(item["orders"]), (row, item)
            assert _money(row[2]) == _money(item["amount"]), (row, item)
            assert _money(row[3]) == _money(item["freight"]), (row, item)

        # ---- 导出（商品经营）----
        sheets, name = _xlsx(client, h, kind="products", **q)
        assert "商品经营" in sheets, f"导出的工作表叫 {list(sheets)}"
        ws = sheets["商品经营"]
        assert name == f"products-report-{anchor}.xlsx", name
        head = _cells(ws, 1)
        assert head[0] == "商品经营" and head[1] == _span_label(today, today), head
        r2 = _cells(ws, 2)
        assert r2[0] == "销售总额" and _money(r2[1]) == _money(p["total_amount"]), (r2, p["total_amount"])
        assert r2[2] == "总件数" and int(r2[3]) == int(p["total_qty"]), (r2, p["total_qty"])
        assert r2[4] == "商品毛利(仅算得出成本的行)", r2
        assert _money(r2[5]) == _money(
            Decimal(str(p["cost_covered_amount"])) - Decimal(str(p["cost_total"]))
        ), (r2[5], p["cost_covered_amount"], p["cost_total"])

        # 商品行：按商品名对齐（顺序由接口决定，导出必须同序）
        item_head_row = next(
            row for row in range(1, ws.max_row + 1) if _cells(ws, row)[:1] == ["商品"]
        )
        assert _cells(ws, item_head_row) == [
            "商品", "件数", "单数", "金额", "参与毛利的金额", "毛利", "货损件数", "货损金额",
        ], _cells(ws, item_head_row)
        got_items = _rows_cols(ws, item_head_row, 8)
        assert len(got_items) == len(p["items"])
        saw_dash = False
        for row, item in zip(got_items, p["items"]):
            assert row[0] == item["product_name"], (row, item)
            assert int(row[1]) == int(item["qty"]), (row, item)
            assert int(row[2]) == int(item["order_count"]), (row, item)
            assert _money(row[3]) == _money(item["amount"]), (row, item)
            assert _money(row[4]) == _money(item["covered_amount"]), (row, item)
            if (item.get("covered_lines") or 0) > 0:
                # 算得出成本 → 必须是一个**数字**，且等于 参与毛利的金额 − 成本
                assert isinstance(row[5], (int, float)), f"毛利应该是数字：{row}"
                assert _money(row[5]) == _money(
                    Decimal(str(item["covered_amount"])) - Decimal(str(item["cost"]))
                ), (row, item)
            else:
                # ⛔ 算不出成本 → 必须写「—」，不许印 0（印 0 会被读成"没成本、毛利全额"）
                assert row[5] == "—", f"算不出成本的行必须写「—」，现在是 {row[5]!r}"
                saw_dash = True
            assert int(row[6]) == int(item["damage_qty"]), (row, item)
            assert _money(row[7]) == _money(item["damage_amount"]), (row, item)
        assert saw_dash, "这一轮没有覆盖到「—」那一格（判据会恒真）"

        # ⑤ 金额列必须是**数**不是文本（Excel 里 SUM 得 0 的老毛病）
        assert isinstance(ws.cell(row=item_head_row + 1, column=4).value, (int, float)), (
            "金额列被写成了文本 —— 用户在 Excel 里 SUM 会得 0"
        )
    finally:
        _cleanup(db_session, order_ids, product_ids)


@pytest.mark.dispatcher
@pytest.mark.fast
@pytest.mark.regression
def test_export_filename_and_series_cover_a_multi_day_window(
    client: TestClient, token_dispatcher: str
) -> None:
    """跨日窗口（近 7 天）：① 文件名与表头是**真实区间**；② 曲线**逐行**与接口一致。

    为什么必须跨日（2026-09-23 反向验证抓到）：单日窗口的曲线只有 1 个点，
    "导出只写第一行"这种缺陷在单日下**看不出来** —— 注入照样全绿。
    """
    h = auth_headers(token_dispatcher)
    today = business_today()
    frm = today - timedelta(days=6)
    s, e = frm.isoformat(), today.isoformat()

    t = client.get(
        "/api/v1/reports/turnover",
        params={"mode": "day", "date": e, "date_from": s, "date_to": e},
        headers=h,
    )
    assert t.status_code == 200, t.text
    series = t.json()["series"]
    assert len(series) >= 2, f"近 7 天的曲线只有 {len(series)} 个点 —— 这条判据会看不出行数问题"

    sheets, name = _xlsx(client, h, kind="turnover", mode="day", date=e, date_from=s, date_to=e)
    assert name == f"turnover-report-{s}_{e}.xlsx", (
        f"区间导出应当按真实区间命名（内容与名字同源）：{name}"
    )
    ws = sheets["营业纵览"]
    assert _cells(ws, 1)[1] == _span_label(frm, today), _cells(ws, 1)

    head_row = next(row for row in range(1, ws.max_row + 1) if _cells(ws, row)[:1] == ["时间"])
    got = _rows_cols(ws, head_row, 4)
    assert len(got) == len(series), f"导出的曲线 {len(got)} 行、接口 {len(series)} 行"
    for row, item in zip(got, series):
        assert row[0] == item["label"], (row, item)
        assert int(row[1]) == int(item["orders"]), (row, item)
        assert _money(row[2]) == _money(item["amount"]), (row, item)
        assert _money(row[3]) == _money(item["freight"]), (row, item)
    # 区间报表的"挂账未收单位TOP"那一段必须跟在曲线后面（不是被曲线吃掉）
    labels = [_cells(ws, r)[:1] for r in range(head_row, ws.max_row + 1)]
    assert ["挂账未收单位TOP"] in labels, f"曲线之后没有「挂账未收单位TOP」那一段：{labels}"
