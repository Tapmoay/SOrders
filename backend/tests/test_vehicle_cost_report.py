"""车辆成本表（FEAT-0012 第二期）：每一台车在这段时间里花了多少钱。

## 为什么要有这个文件（2026-10-04）

折旧接进利润表之后，下一个问题一定是「哪台车在烧钱」—— 而三笔成本本来散在三个地方：
折旧在车辆台账、这台车的开销在收支页、配送成本在运费结算页。按车摆在一起的入口一个都没有。
`GET /reports/vehicle-cost` 就是那一处，且**一分钱都不自己算**：三笔全部调既有唯一实现
（`services/vehicle_depreciation.py` / `expenses.vehicle_id` / `money_contract.pay_for_order`）。

这个文件钉六件事：
1. 恒等式当场成立：逐车 `total_cost == 折旧 + 开销 + 配送成本`，且合计 == 逐车相加；
2. 台账四格录全的车，折旧**当场**进表（按自然月天数摊，只四舍五入一次）；
3. 缺购置信息的车：折旧是 0 而**月额是 null**（算不出来 ≠ 0），缺哪几格逐条列出；
4. 开销按 `exp_date` 落窗口、按分类聚合（金额降序），窗口外的一分钱都不进；
5. 配送成本与运费结算/利润表**同源**（同一个 `pay_for_order`）；
6. ⛔ 这张表**没有收入**：订单上没有「哪台车拉的」这个事实，硬摊就是编一个比例。
"""

from __future__ import annotations

import itertools
import json
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from starlette.testclient import TestClient

from tests.conftest import auth_headers

PROBE = "车辆成本探针"
_SEQ = itertools.count(1)


def _get(client: TestClient, h: dict[str, str], **params) -> dict:
    q = "&".join(f"{k}={v}" for k, v in params.items() if v is not None)
    r = client.get(f"/api/v1/reports/vehicle-cost?{q}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _dec(d: dict, key: str) -> Decimal:
    return Decimal(str(d[key]))


def _row(data: dict, vehicle_id: int) -> dict:
    """按车找那一行（⛔ 不许按顺序取：车的顺序变了测试就假红/假绿）。"""
    for row in data["per_vehicle"]:
        if row["vehicle_id"] == vehicle_id:
            return row
    raise AssertionError(f"表里没有这台车：{vehicle_id}")


def _mk_vehicle(client: TestClient, h: dict[str, str], plate: str, **fields) -> int:
    r = client.post("/api/v1/vehicles", json={"plate_no": plate, "vehicle_type": "small", **fields}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _mk_expense(db_session, *, day: date, category: str, amount: str, vehicle_id: int) -> None:
    """直接落一条挂在这台车上的开销（本文件只读报表，⛔ 不走「保存后自动生成资金流水」那条写链路）。"""
    from app.models.expense import Expense

    db_session.add(
        Expense(
            exp_date=day,
            category=category,
            amount=Decimal(amount),
            vehicle_id=vehicle_id,
            note=PROBE,
        )
    )
    db_session.commit()


@pytest.fixture
def cleanup_vehicles(db_session):
    """跑完把本文件造的车与开销行收拾干净（测试库整个会话共用一个文件）。"""
    from app.models import Vehicle
    from app.models.expense import Expense

    before = {v.id for v in db_session.scalars(select(Vehicle)).all()}
    yield
    db_session.rollback()
    for exp in db_session.scalars(select(Expense).where(Expense.note == PROBE)).all():
        db_session.delete(exp)
    for v in db_session.scalars(select(Vehicle)).all():
        if v.id not in before:
            db_session.delete(v)
    db_session.commit()


def test_恒等式与口径说明(client, token_dispatcher, cleanup_vehicles):
    """逐车三笔相加 == 合计、合计 == 逐车相加；口径说明逐条写着（⛔ 无 markdown 星号）。"""
    h = auth_headers(token_dispatcher)
    covered = _mk_vehicle(
        client, h, "测本00011",
        purchase_price="120000", purchase_date="2025-09-16",
        useful_life_years="5", residual_rate="0.05",
    )
    uncovered = _mk_vehicle(client, h, "测本00012")
    d = _get(client, h, mode="month", date="2026-09-30")
    per = d["per_vehicle"]
    assert d["vehicle_count"] == len(per) >= 2
    assert _row(d, covered)["depreciation_covered"] is True
    assert _row(d, uncovered)["depreciation_covered"] is False
    for row in per:
        assert _dec(row, "total_cost") == (
            _dec(row, "depreciation") + _dec(row, "expense_total") + _dec(row, "delivery_cost")
        ), row["plate_no"]
    assert _dec(d, "total_cost") == sum((_dec(r, "total_cost") for r in per), Decimal("0"))
    assert _dec(d, "total_cost") == (
        _dec(d, "depreciation_total") + _dec(d, "expense_total") + _dec(d, "delivery_cost_total")
    )
    assert len(d["notes"]) >= 5
    assert not any("**" in n for n in d["notes"]), "口径说明会原样进手机与导出的表"
    # ⛔ 这张表没有收入：订单上没有「哪台车拉的」这个事实（硬摊就是编一个比例）
    keys = set(d) | {k for row in per for k in row}
    assert not [k for k in keys if "revenue" in k or "income" in k]


def test_新车折旧当场进表(client, token_dispatcher, cleanup_vehicles):
    """四格录全 -> 折旧当场进表；整月正好是月额（120000×95%÷60 = 1900.00）。"""
    h = auth_headers(token_dispatcher)
    before = _dec(_get(client, h, mode="month", date="2026-09-30"), "depreciation_total")
    vid = _mk_vehicle(
        client, h, "测本00001",
        purchase_price="120000", purchase_date="2025-09-16",
        useful_life_years="5", residual_rate="0.05",
    )
    d = _get(client, h, mode="month", date="2026-09-30")
    row = _row(d, vid)
    assert row["depreciation_covered"] is True
    assert _dec(row, "monthly_depreciation") == Decimal("1900.00")
    assert _dec(row, "depreciation") == Decimal("1900.00"), "整月就是一个月额"
    assert _dec(d, "depreciation_total") - before == Decimal("1900.00")
    assert d["covered_count"] >= 1, "有购置信息的车要算进「算得出折旧」那一格"
    assert vid not in [r["vehicle_id"] for r in d["per_vehicle"] if not r["depreciation_covered"]]


def test_开销按分类进表且窗口外不进(client, token_dispatcher, db_session, cleanup_vehicles):
    """开销按 `exp_date` 落窗口、按分类聚合（金额降序）；窗口外的一分钱都不进。"""
    h = auth_headers(token_dispatcher)
    vid = _mk_vehicle(client, h, "测本00002")
    _mk_expense(db_session, day=date(2026, 9, 20), category="停车", amount="900.00", vehicle_id=vid)
    _mk_expense(db_session, day=date(2026, 9, 21), category="过路", amount="80.00", vehicle_id=vid)
    _mk_expense(db_session, day=date(2026, 8, 31), category="保险", amount="1500.00", vehicle_id=vid)
    row = _row(_get(client, h, mode="month", date="2026-09-30"), vid)
    assert [(x["category"], _dec(x, "amount")) for x in row["expenses"]] == [
        ("停车", Decimal("900.00")),
        ("过路", Decimal("80.00")),
    ]
    assert _dec(row, "expense_total") == Decimal("980.00"), "8-31 那条是窗口外的"


def test_缺购置信息的车算不出来不是0(client, token_dispatcher, cleanup_vehicles):
    """缺格 -> 折旧 0 但月额是 null、缺哪几格逐条列出（⛔ 不猜一个数出来）。"""
    h = auth_headers(token_dispatcher)
    vid = _mk_vehicle(client, h, "测本00003", purchase_price="60000")
    row = _row(_get(client, h, mode="month", date="2026-09-30"), vid)
    assert row["depreciation_covered"] is False
    assert row["depreciation_uncovered_reasons"] == ["没录购置日期", "没录使用年限"]
    assert row["monthly_depreciation"] is None, "算不出来是 null，⛔ 不是 0"
    assert _dec(row, "depreciation") == Decimal("0")


def test_配送成本与运费结算同源(client, token_dispatcher, db_session, users, cleanup_vehicles):
    """挂在这台车上的司机，在这一天里按单应付多少 —— 与运费结算页同一个函数算出来的。"""
    from app.models import Order, OrderProduct
    from app.models.enums import OrderStatus
    from app.core.business_time import business_today, utc_now_naive
    from app.services.money_contract import has_per_order_pay, pay_for_order
    from app.services.reports.loader import load_delivered

    h = auth_headers(token_dispatcher)
    vid = _mk_vehicle(client, h, "测本00004", driver_id=users["driver"].id)
    day = business_today()
    order = Order(
        order_no=f"VCOST{day.strftime('%m%d')}-{next(_SEQ)}",
        status=OrderStatus.DELIVERED,
        order_date=day,
        shipper_id=users["shipper"].id,
        driver_id=users["driver"].id,
        delivered_at=utc_now_naive(),
        dispatched_at=utc_now_naive(),
        freight_fee=Decimal("20.00"),
        driver_billing_mode_snapshot="PIECE",
        delivery_description=PROBE,
    )
    db_session.add(order)
    db_session.flush()
    db_session.add(
        OrderProduct(
            order_id=order.id,
            product_name_snapshot=PROBE,
            quantity=2,
            unit_price=Decimal("50.00"),
            line_total=Decimal("100.00"),
            cost_price_snapshot=Decimal("30.00"),
        )
    )
    db_session.commit()
    try:
        row = _row(_get(client, h, mode="day", date=day.isoformat()), vid)
        mine = [
            o
            for o in load_delivered(db_session, span=(day, day))
            if o.driver_id == users["driver"].id and has_per_order_pay(o)
        ]
        expected = sum((pay_for_order(o).total for o in mine), Decimal("0"))
        assert expected > 0
        assert _dec(row, "delivery_cost") == expected
    finally:
        db_session.query(OrderProduct).filter(OrderProduct.order_id == order.id).delete()
        db_session.query(Order).filter(Order.id == order.id).delete()
        db_session.commit()


def test_没有派单权限的看不到这张表(client, token_shipper):
    """报表是全店口径（没有「只看自己那份」的版本）—— 不是派单员就是 403。"""
    r = client.get("/api/v1/reports/vehicle-cost?mode=month&date=2026-09-30", headers=auth_headers(token_shipper))
    assert r.status_code == 403


def test_导出kind车辆成本有那张表(client, token_dispatcher):
    """导出分支：kind=vehicle-cost 出「车辆成本」那张表（表头 + 明细 + 口径说明）。"""
    h = auth_headers(token_dispatcher)
    r = client.get("/api/v1/reports/export?kind=vehicle-cost&mode=month&date=2026-09-30", headers=h)
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert "车辆成本" in wb.sheetnames
    ws = wb["车辆成本"]
    cells = [c.value for row in ws.iter_rows() for c in row]
    for want in ("车辆成本", "车牌", "挂靠司机", "车辆折旧", "这台车的开销", "配送成本", "成本合计", "口径说明"):
        assert want in cells, want
