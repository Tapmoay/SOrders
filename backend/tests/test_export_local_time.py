"""导出 xlsx 里的**时间列**必须是"给人看的当地时刻"，而且同一张 sheet 里口径一致。

## 抓到的形状（2026-09-24 第 20 轮并行渗透 D3-F3）
`/reports/export?kind=audit` 那张 sheet 上有**两个时间列**，而它们原来印的是**两个口径**：

| 列 | 原来 | 实测（同一件事，相差 1 毫秒的两行） |
| --- | --- | --- |
| 异常单的「解决时间」 | `exception_resolved_at.isoformat()` = **UTC naive** | `2026-09-21T11:50:01.206089` |
| 下方「敏感操作日志」的「时间」 | `local_stamp`（第 19 轮修的）= **当地** | `2026-09-21 19:50` |

同一张表上一件事差 8 小时，看的人第一反应是"导出的窗口错了"。
另外 `stats_export` 的「实际送达」「约定送达前」也是原始 UTC（同族）。
口径只有一处：`business_time.local_stamp`（要带年份就显式给 `fmt`）。

⚠️ 这条判据同时也钉住"**不要**为了省事把 `T`/微秒带出去"——Excel 认不出那种字符串是时间，
按日期排序/筛选会失效（本仓库的历史产物就是 `2026-09-02` 这样的 `inlineStr`）。
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import auth_headers


def _cells(ws, row: int) -> list:
    return [c.value for c in ws[row]]


def test_审计导出的两个时间列都是当地时刻(client, db_session, users, token_dispatcher):
    from app.models import Order
    from app.models.enums import OrderStatus, OperationAction
    from app.models.operation_log import OperationLog

    h = auth_headers(token_dispatcher)
    # 2026-09-21 11:50 UTC = 当地（东八区）19:50 —— 两个列必须都印 19:50
    utc_moment = datetime(2026, 9, 21, 11, 50, 1, tzinfo=timezone.utc).replace(tzinfo=None)
    o = Order(
        order_no="SO-EXPORT-TZ-1",
        shipper_id=users["shipper"].id,
        driver_id=users["driver"].id,
        status=OrderStatus.DELIVERED,
        order_date=date(2026, 9, 21),
        is_exception=True,
        exception_reason="导出时间口径探针",
        exception_resolved_at=utc_moment,
    )
    db_session.add(o)
    db_session.flush()
    db_session.add(
        OperationLog(
            operator_id=users["dispatcher"].id,
            order_id=o.id,
            action=OperationAction.ORDER_UPDATE,
            change_content="导出时间口径探针",
            created_at=utc_moment,
        )
    )
    db_session.commit()

    r = client.get(
        "/api/v1/reports/export?kind=audit&mode=day&date=2026-09-21", headers=h
    )
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active

    flat = [
        str(c.value)
        for row in ws.iter_rows()
        for c in row
        if c.value not in (None, "")
    ]
    stamps = [v for v in flat if "2026-09-21" in v and (":" in v)]
    assert stamps, f"导出里没找到任何时间戳：{flat[-12:]}"
    for v in stamps:
        assert "19:50" in v, (
            f"导出里的时间不是当地时刻：{v!r}（UTC 的 11:50 应当印成 19:50）—— "
            "同一张 sheet 里两个时间列口径不一致，看的人会以为窗口错了"
        )
        assert "T" not in v, f"还带着 ISO 的 `T` 与微秒：{v!r}（Excel 认不出它是时间，排序会失效）"


def test_导出里的等号开头文本是文本不是活公式(client, db_session, users, token_dispatcher):
    """公式注入（2026-09-24 第 20 轮 D7-1）：以 `=` 开头的文本不许在产物里变成 `<f>` 公式。

    用**审计块**验：那一块印的是 `operation_logs.change_content`（"内容"列），
    而它是**用户可控**的（每一次改动都会把值写进去），也是 D7-1 点名的入口之一。
    """
    from app.models.enums import OperationAction
    from app.models.operation_log import OperationLog

    h = auth_headers(token_dispatcher)
    evil = '=HYPERLINK("http://example.invalid/","点我")'
    db_session.add(
        OperationLog(
            operator_id=users["dispatcher"].id,
            order_id=None,
            action=OperationAction.ORDER_UPDATE,
            change_content=evil,
            created_at=datetime(2026, 9, 21, 11, 50, 1, tzinfo=timezone.utc).replace(tzinfo=None),
        )
    )
    db_session.commit()

    r = client.get("/api/v1/reports/export?kind=audit&mode=day&date=2026-09-21", headers=h)
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    found = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value == evil:
                    found = True
                    assert c.data_type == "s", (
                        f"格子 {c.coordinate} 的类型是 {c.data_type!r} —— "
                        "以 `=` 开头的用户文本被写成了**活公式**（Excel 打开即求值）"
                    )
    assert found, "探针数据没出现在导出里（这条判据没验到东西）"
