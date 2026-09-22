"""第二轮外部完整检查报告里**账本/报表导出内容**一节的判据。

报告是**仓库之外**产出的（条目带编号与证据行号），本文件只钉"能被一条断言证明、
且退回旧写法就立刻变红"的行为；每条都做过注入式反向验证（把修复退回 → 红；改回 → 绿）。

| 条目 | 缺陷形状 | 判据 |
|---|---|---|
| **R2-1** | 账本导出的文本列以 `=` 开头 → openpyxl 写出 `<f>` **公式节点** | 产物 XML 里没有公式节点，且值逐字不变 |
| **R2-6(exp)** | `write_excel` 在循环里 `db.get(Order, ...)` → 每行一条 SQL | 一次导出只发**一条** `FROM orders`，且每行仍是它自己的订单说明 |
| **R2-3(exp)** | 司机绩效「待结运费」写成字符串 → Excel 里 `SUM` 得 0 | 计件司机那一格是**数字**节点；工资制仍是文字"工资制" |
| **R2-7(exp)** | 客户经营导出把整张 `ledgers` 读进内存，再在 Python 里按日期过滤 | 条件下推到 SQL，且逐行结果与"全表读 + Python 过滤"**完全一致** |
| **R2-4** | `turnover`/`products` 文件名写锚点日、内容是整月 | 文件名与真实取数区间**同源** |

判定一律读产物里的**原始 XML**（`xl/worksheets/sheet1.xml`），不看 openpyxl 的对象属性 ——
缺陷现场恰恰是"对象上是个字符串、写出来却是公式节点"（`cell.data_type` 反映的是我们设进去的值，
不是写出去的节点）。R2-1 的反面证据见本文件第一条用例：`ws.append(["=1+1"])` 的产物里就是 `<f>1+1</f>`。
"""
from __future__ import annotations

import re
import warnings
import zipfile
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from tests.conftest import auth_headers

# 一份"以危险字符开头"的备注：CSV 注入时代的标准探针形状
DANGER_NOTE = '=HYPERLINK("http://evil.example/x","点我")'
DANGER_ORDER_NO = "=cmd|'/C calc'!A0"
DANGER_DESC = "=1+1"
PROBE = "R2导出探针"


# ---------------------------------------------------------------- 产物判据的工具
def _sheet_xml(content: bytes, sheet_index: int = 1) -> str:
    """从 xlsx 产物里读**原始 XML**：公式还是文本，只有这里说了算。"""
    with zipfile.ZipFile(BytesIO(content)) as z:
        return z.read(f"xl/worksheets/sheet{sheet_index}.xml").decode("utf-8")


def _cell_xml(xml: str, ref: str) -> str:
    """取一个格子的节点原文（`<c r="A1" ...>...</c>`）。"""
    m = re.search(rf'<c r="{ref}"(?:[^>]*/>|[^>]*>.*?</c>)', xml, re.S)
    assert m, f"产物里没有 {ref} 这个格子"
    return m.group(0)


def _export(client, h, **params) -> tuple[str, bytes]:
    """导出一份报表，返回（文件名, xlsx 字节）。"""
    from urllib.parse import urlencode

    r = client.get(f"/api/v1/reports/export?{urlencode(params)}", headers=h)
    assert r.status_code == 200, r.text
    disp = r.headers.get("content-disposition") or ""
    return disp.split("filename=")[-1].strip('"'), r.content


def _sheet(content: bytes, title: str):
    return load_workbook(BytesIO(content))[title]


def _data_rows(ws, min_row: int = 3) -> list[list]:
    """表里的数据行（到第一个空行为止）——用于"逐行一致"的比较。"""
    out: list[list] = []
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if row[0] in (None, ""):
            break
        out.append(list(row))
    return out


def _capture_sql(call):
    """跑一次 [call]，返回（结果, 它执行过的所有 ORM 语句文本）。

    ⚠️ 走 `Session.do_orm_execute`（**类级**监听），不是引擎的 `before_cursor_execute`：
    引擎级监听只对它**注册之后新建的连接**生效，而测试里那条连接早被 fixture 建好了 ——
    本机实测同一个导出请求：引擎监听捕获 **0** 条语句，`do_orm_execute` 捕获 2 条
    （账本一条 + 挂账汇总一条）。"到底发了几条 SQL"必须真的看得见，
    否则 N+1 那条判据就是恒绿的摆设。
    """
    from sqlalchemy import event
    from sqlalchemy.orm import Session as SASession

    seen: list[str] = []

    def _rec(state):
        try:
            with warnings.catch_warnings():
                # 字面量化会把别的语句里的 `= NULL` 渲染出来并告警；这里只要 SQL 文本
                warnings.simplefilter("ignore")
                seen.append(str(state.statement.compile(compile_kwargs={"literal_binds": True})))
        except Exception:  # noqa: BLE001 - 有些类型没法字面量化，退回不带绑定的形态
            seen.append(str(state.statement))

    event.listen(SASession, "do_orm_execute", _rec)
    try:
        result = call()
    finally:
        event.remove(SASession, "do_orm_execute", _rec)
    return result, seen


def _purge(db_session, *, ledger_ids=(), order_ids=()) -> None:
    """删掉本文件造的探针数据。

    这个测试库是 **session 级共享**的（一个 worker 一个库文件），留下的行会进别的用例的合计；
    先删账本行与订单行再删订单本身（`PRAGMA foreign_keys=ON`，外键是反的会报错）。
    """
    from app.models import Ledger, Order, OrderProduct

    if ledger_ids:
        db_session.query(Ledger).filter(Ledger.id.in_(list(ledger_ids))).delete(synchronize_session=False)
    if order_ids:
        db_session.query(OrderProduct).filter(OrderProduct.order_id.in_(list(order_ids))).delete(
            synchronize_session=False
        )
        db_session.query(Order).filter(Order.id.in_(list(order_ids))).delete(synchronize_session=False)
    db_session.commit()


def _mk_orders_and_ledgers(db_session, shipper_id: int, *, tag: str, orders: list[tuple[str, str]], n_ledger: int):
    """造 [orders]（order_no, delivery_description）+ [n_ledger] 行账本（轮流挂到这些单上）。"""
    from app.models import Ledger, Order
    from app.models.enums import LedgerSource, OrderStatus

    rows_o = [
        Order(
            order_no=no,
            status=OrderStatus.DELIVERED,
            order_date=date(2026, 3, 10),
            shipper_id=shipper_id,
            delivery_description=desc,
        )
        for no, desc in orders
    ]
    db_session.add_all(rows_o)
    db_session.flush()
    rows_l = [
        Ledger(
            shipper_id=shipper_id,
            entry_date=date(2026, 3, 10),
            product_name=f"{tag}{i}",
            quantity=1,
            unit_price=Decimal("10.00"),
            total=Decimal("10.00"),
            order_id=rows_o[i % len(rows_o)].id,
            source=LedgerSource.MANUAL,
            note=f"{tag}{i}",
        )
        for i in range(n_ledger)
    ]
    db_session.add_all(rows_l)
    db_session.commit()
    return rows_o, rows_l


# ---------------------------------------------------------------- ① R2-1 公式注入
def test_ledger_excel_text_cells_are_text_nodes_never_formulas(db_session, users, tmp_path):
    """以 `=` 开头的货主名/订单说明/商品名/订单号/备注，在产物里必须是**文本节点**。

    ⛔ 反面证据（同一台机器同一条路径，缺陷现场）：原来这里是 `ws.append([...])`，
    `=1+1` 写出来是 `<c r="A1"><f>1+1</f><v></v></c>` —— 活公式。
    本文件第一条断言就是"产物里一个 `<f>` 都没有"（这份表里本来就没有任何合法公式）。
    """
    from app.services.ledger_export import write_excel
    from tests.conftest import get_test_session_factory

    # n_ledger=2 → 第一行挂在"危险"那张单上，第二行挂在中文那张单上
    orders, ledgers = _mk_orders_and_ledgers(
        db_session,
        users["shipper"].id,
        tag=PROBE,
        orders=[(DANGER_ORDER_NO, DANGER_DESC), ("SO20260310001", "给张老板的一车货")],
        n_ledger=2,
    )
    danger_l, normal_l = ledgers[0], ledgers[1]
    danger_l.note = DANGER_NOTE
    danger_l.product_name = DANGER_DESC
    normal_l.note = "备注：早上送到，别压坏"
    normal_l.product_name = "土豆"
    db_session.commit()

    path = tmp_path / "ledger.xlsx"
    session = get_test_session_factory()()
    try:
        # 用**新会话**导出：与生产一致（导出跑在自己的 Session 上），身份映射也是空的
        write_excel(session, path, [danger_l, normal_l], "=危险货主", date(2026, 3, 1), date(2026, 3, 31))
    finally:
        session.close()

    xml = _sheet_xml(path.read_bytes())
    assert "<f>" not in xml and "<f " not in xml, (
        "产物里出现了公式节点 —— 打开这份 xlsx 的人 Excel 会去执行格子里的内容"
    )

    # 文本列：日期(A)/订单说明(B)/商品(C)/订单号(H)/来源(I)/备注(J)
    for ref in ("B3", "C3", "H3", "J3"):
        cell = _cell_xml(xml, ref)
        assert 't="inlineStr"' in cell and "<is>" in cell, f"{ref} 不是文本节点：{cell}"

    ws = _sheet(path.read_bytes(), "账本")
    assert ws["B1"].value == "=危险货主", "货主名被改动了"
    assert ws["B3"].value == DANGER_DESC, f"订单说明没逐字保留：{ws['B3'].value!r}"
    assert ws["C3"].value == DANGER_DESC
    assert ws["H3"].value == DANGER_ORDER_NO, "订单号没逐字保留"
    assert ws["J3"].value == DANGER_NOTE, "备注没逐字保留"
    # 正常中文一个字都不许变（没有"加前导引号/空格"这种妥协）
    assert [ws.cell(4, c).value for c in range(1, 11)] == [
        "2026-03-10", "给张老板的一车货", "土豆", 1, 10.0, 10.0,
        orders[1].id, "SO20260310001", "manual", "备注：早上送到，别压坏",
    ], "正常中文行的内容变了"

    _purge(db_session, ledger_ids=[danger_l.id, normal_l.id], order_ids=[o.id for o in orders])


def test_ledger_excel_number_columns_stay_numbers(db_session, users, tmp_path):
    """数量/单价/总价必须仍是**数值**格子（Excel 里要能求和）。

    这条是给 ① 那个"一律转成文本"的偷懒修法准备的：把整行都当文本会让账本再也算不出合计，
    与 R2-3 修的是同一件事的两半。
    """
    from app.models import Ledger
    from app.services.ledger_export import write_excel
    from tests.conftest import get_test_session_factory

    orders, ledgers = _mk_orders_and_ledgers(
        db_session, users["shipper"].id, tag=PROBE, orders=[("SO20260310NUM", "数值列探针")], n_ledger=1
    )
    ledgers[0].quantity = 3
    ledgers[0].unit_price = Decimal("12.50")
    ledgers[0].total = Decimal("37.50")
    db_session.commit()

    path = tmp_path / "num.xlsx"
    session = get_test_session_factory()()
    try:
        write_excel(session, path, ledgers, "货主", date(2026, 3, 1), date(2026, 3, 31))
    finally:
        session.close()

    xml = _sheet_xml(path.read_bytes())
    for ref, want in (("D3", 3), ("E3", 12.5), ("F3", 37.5)):
        cell = _cell_xml(xml, ref)
        assert 't="inlineStr"' not in cell, f"{ref} 被写成了文本：{cell}"
        assert "<is>" not in cell, f"{ref} 被写成了文本：{cell}"
        assert f"<v>{want}</v>" in cell, f"{ref} 不是数值节点：{cell}"

    ws = _sheet(path.read_bytes(), "账本")
    assert ws["D3"].value == 3 and ws["E3"].value == 12.5 and ws["F3"].value == 37.5

    _purge(db_session, ledger_ids=[r.id for r in ledgers], order_ids=[o.id for o in orders])


# ---------------------------------------------------------------- ② R2-6 每行一次 SELECT
def test_ledger_excel_queries_orders_once_not_per_row(db_session, users, tmp_path):
    """5 行账本（挂 3 张单）只许发**一条** `SELECT ... FROM orders`，且每行文本仍是自己的单。

    ⚠️ 用**新会话**导出是判据的一部分：`db.get` 命中身份映射时不会发 SQL，
    拿"测试里刚创建过这些单"的会话去测会把这处 N+1 测没。生产里导出跑在自己的 Session 上。
    """
    from app.models import Ledger
    from app.services.ledger_export import write_excel
    from tests.conftest import get_test_session_factory

    orders, ledgers = _mk_orders_and_ledgers(
        db_session,
        users["shipper"].id,
        tag=PROBE,
        orders=[("R2N0", "第0单说明"), ("R2N1", "第1单说明"), ("R2N2", "第2单说明")],
        n_ledger=5,
    )
    ids = [r.id for r in ledgers]

    session = get_test_session_factory()()
    try:
        rows = list(session.scalars(select(Ledger).where(Ledger.id.in_(ids)).order_by(Ledger.id)))
        assert len(rows) == 5
        path = tmp_path / "n1.xlsx"
        _, sqls = _capture_sql(
            lambda: write_excel(session, path, rows, "货主", date(2026, 3, 1), date(2026, 3, 31))
        )
    finally:
        session.close()

    order_selects = [q for q in sqls if "from orders" in q.lower()]
    assert len(order_selects) == 1, (
        f"订单说明的查询发了 {len(order_selects)} 条（N+1）：一份 500 行的账本 = 501 条 SQL。"
        f"实际语句：{order_selects}"
    )
    assert " in (" in order_selects[0].lower(), f"没有批量化：{order_selects[0]}"

    # 顺序与内容严格不变：第 i 行必须还是**它自己那张单**的说明与单号
    ws = _sheet(path.read_bytes(), "账本")
    got = [(ws.cell(i, 2).value, ws.cell(i, 8).value) for i in range(3, 8)]
    assert got == [(f"第{i % 3}单说明", f"R2N{i % 3}") for i in range(5)], (
        f"批量查询把行的订单说明/单号串了：{got}"
    )

    _purge(db_session, ledger_ids=ids, order_ids=[o.id for o in orders])


# ---------------------------------------------------------------- ③ R2-3 待结运费是文本
def test_drivers_export_freight_owed_is_numeric_but_salary_stays_text(
    client, token_dispatcher, db_session, users
):
    """计件司机的「待结运费」在 xlsx 里必须是**数值**；工资制那一格必须仍是文字"工资制"。

    - 数值：整列是文本时，用户在 Excel 里选这一列 `SUM` 得 0（R2-3）；
    - 文字：印 0 会被读成"这个月一分钱都不用付给他"（这是**刻意**的设计，页面同款）。
    """
    from app.core.business_time import business_today, utc_now_naive
    from app.models import Order, OrderProduct
    from app.models.enums import OrderStatus

    driver = users["driver"]
    # 显式给定前置状态：不依赖别的用例有没有给这个司机挂过规则（挂了规则时 snapshot_mode 看规则）
    saved = (driver.billing_mode, driver.driver_rule_id)
    driver.billing_mode = "PIECE"
    driver.driver_rule_id = None
    db_session.commit()

    day = business_today()
    order = Order(
        order_no="R2DRIVER1",
        status=OrderStatus.DELIVERED,
        order_date=day,
        shipper_id=users["shipper"].id,
        driver_id=driver.id,
        delivered_at=utc_now_naive(),
        dispatched_at=utc_now_naive(),
        freight_fee=Decimal("100.00"),
        driver_billing_mode_snapshot="PIECE",
        delivery_description=PROBE,
    )
    db_session.add(order)
    db_session.flush()
    db_session.add(
        OrderProduct(
            order_id=order.id,
            product_name_snapshot=PROBE,
            quantity=1,
            unit_price=Decimal("100.00"),
            line_total=Decimal("100.00"),
        )
    )
    db_session.commit()

    h = auth_headers(token_dispatcher)
    window = {"mode": "day", "date": day.isoformat(), "date_from": day.isoformat(), "date_to": day.isoformat()}
    name = driver.full_name or driver.phone

    def _driver_row():
        _fn, content = _export(client, h, kind="drivers", **window)
        ws = _sheet(content, "司机绩效")
        for idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            if row[0] == name:
                return content, idx, row
        raise AssertionError(f"导出里没有这个司机：{name}")

    try:
        # 页面口径：这一格在 JSON 里**就是字符串**（`DriverPerformanceRow.freight_owed: str | None`）——
        # 导出那一条不许反过来把这个接口改成数字（页面/App 按字符串格式化它）
        page = client.get(
            f"/api/v1/stats/driver-performance?date_from={day}&date_to={day}", headers=h
        ).json()
        mine = [d for d in page["drivers"] if d["driver_id"] == driver.id]
        assert mine, f"绩效页里没有这个司机：{page}"
        assert mine[0]["billing_mode"] == "PIECE", mine[0]
        owed = mine[0]["freight_owed"]
        assert isinstance(owed, str), f"页面那一格的口径被改了（不是字符串）：{owed!r}"

        content, idx, row = _driver_row()
        assert row[5] == "PIECE", f"计费方式列变了：{row}"
        cell = _cell_xml(_sheet_xml(content), f"G{idx}")
        assert 't="inlineStr"' not in cell and "<is>" not in cell and "<v>" in cell, (
            f"待结运费写成了文本节点 → Excel 里这一列 SUM 得 0：{cell}"
        )
        want = float(Decimal(owed).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        assert row[6] == pytest.approx(want), f"导出的钱与页面不同源：{row[6]} vs {owed}"

        # 工资制：同一列必须是文字（刻意设计，页面同款）
        driver.billing_mode = "SALARY"
        db_session.commit()
        content2, idx2, row2 = _driver_row()
        assert row2[5] == "SALARY", f"计费方式列变了：{row2}"
        assert row2[6] == "工资制", f"工资制司机这一格印成了 {row2[6]!r}（会被读成一分钱都不用付）"
        cell2 = _cell_xml(_sheet_xml(content2), f"G{idx2}")
        # ⚠️ 中文在 XML 里是数字字符引用（`&#24037;`），所以只判节点类型 + 读回来的值
        assert 't="inlineStr"' in cell2 and "<is>" in cell2, f"文字格子的节点不对：{cell2}"
    finally:
        driver.billing_mode, driver.driver_rule_id = saved
        db_session.commit()
        _purge(db_session, order_ids=[order.id])


# ---------------------------------------------------------------- ④ R2-7 日期条件下推
def _buckets_as_pre_fix(db_session, s: date, e: date) -> list[list]:
    """把**修复前**的算法原样写一遍：全表读 → 在 Python 里按日期过滤 → 同一套分桶/排序。

    这是 ④ 的"逐行一致"基准：判据不是"改完之后看着对"，而是"与改动前的输出完全相同"。
    """
    from app.api.v1.reports import _money
    from app.models import Ledger, User
    from app.services.ledger_scope import visible_ledger_select

    rows = list(db_session.scalars(visible_ledger_select()))
    user_cache: dict[int, User | None] = {}
    shipper_buckets: dict[int, dict] = {}
    member_buckets: dict[int, dict] = {}
    temp_bucket: dict[str, dict] = {}
    for r in rows:
        if r.entry_date < s or r.entry_date > e:
            continue
        if r.shipper_id is not None:
            if r.shipper_id not in user_cache:
                user_cache[r.shipper_id] = db_session.get(User, r.shipper_id)
            u = user_cache[r.shipper_id]
            b = (
                member_buckets if (u is not None and getattr(u, "is_member", False)) else shipper_buckets
            ).setdefault(
                r.shipper_id,
                {
                    "name": (u.full_name or u.phone or f"货主#{r.shipper_id}") if u else f"货主#{r.shipper_id}",
                    "count": 0,
                    "total": Decimal("0"),
                },
            )
        else:
            nm = (r.temp_shipper_name or "").strip() or "临时货主"
            b = temp_bucket.setdefault(nm, {"name": nm, "count": 0, "total": Decimal("0")})
        b["count"] += 1
        b["total"] += r.total or Decimal("0")

    out = [
        ["货主", b["name"], b["count"], _money(b["total"])]
        for b in sorted(shipper_buckets.values(), key=lambda x: -x["total"])
    ]
    out += [
        ["临时货主", b["name"], b["count"], _money(b["total"])]
        for b in sorted(temp_bucket.values(), key=lambda x: -x["total"])
    ]
    out += [
        ["批发商", b["name"], b["count"], _money(b["total"])]
        for b in sorted(member_buckets.values(), key=lambda x: -x["total"])
    ]
    return out


def test_customers_export_pushes_dates_into_sql_and_keeps_rows_identical(
    client, token_dispatcher, db_session, users
):
    """客户经营导出：日期条件下推到 SQL 之后，**逐行**结果必须与"全表读 + Python 过滤"一模一样。

    窗口取 2026-03（全仓库的用例都不碰这个月），所以断言可以精确到"我造的那几条"。
    """
    from app.models import Ledger
    from app.models.enums import LedgerSource

    s, e = date(2026, 3, 10), date(2026, 3, 20)
    tag = f"{PROBE}客户经营"
    probes: list[Ledger] = []

    def _add(entry_date: date, amount: str, **kw) -> None:
        row = Ledger(
            entry_date=entry_date,
            product_name=tag,
            quantity=1,
            unit_price=Decimal(amount),
            total=Decimal(amount),
            source=LedgerSource.MANUAL,
            note=tag,
            **kw,
        )
        db_session.add(row)
        probes.append(row)

    _add(date(2026, 3, 9), "111.00", shipper_id=users["shipper"].id)   # 区间外（早一天）
    _add(s, "222.00", shipper_id=users["shipper"].id)                  # 左边界，含
    _add(e, "333.00", shipper_id=users["shipper"].id)                  # 右边界，含
    _add(date(2026, 3, 21), "444.00", shipper_id=users["shipper"].id)  # 区间外（晚一天）
    _add(date(2026, 3, 8), "555.00", temp_shipper_name=f"{tag}区间外临时户")  # 只在区间外
    db_session.commit()

    h = auth_headers(token_dispatcher)
    params = {"kind": "customers", "mode": "day", "date": s.isoformat(), "date_from": s.isoformat(), "date_to": e.isoformat()}
    try:
        (fn, content), sqls = _capture_sql(lambda: _export(client, h, **params))
        assert fn == f"customers-report-{s}_{e}.xlsx", f"文件名与真实区间不符：{fn}"

        # 下推的证据：**取行**的那条查询（`select ledgers.id, …`）只有一条，且带着 entry_date 的上下界
        # （`DELETE FROM ledgers` 是保留任务在顺带清理三年以上的老账，不算在"取数"里）
        #
        # ⚠️ 2026-09-20：这里从"只有一条 ledgers 查询"放宽成"**取行的**只有一条 + 总数有上界"。
        #    原因是加了退货之后，`order_money.money_map` 会对同一批订单再发一条
        #    `select ledgers.order_id, sum(...) ... group by order_id` 的**聚合**查询
        #    （退货红冲金额）。它不随结果行数增长（一次算完一页），所以不是这条测试要拦的 N+1；
        #    真正的判据是"取行的那条不许变多"。
        ledger_selects = [
            q for q in sqls if q.lstrip().lower().startswith("select") and "from ledgers" in q.lower()
        ]
        ledger_row_selects = [q for q in ledger_selects if "ledgers.id" in q.lower()]
        assert len(ledger_row_selects) == 1, f"取行的账本查询不是一条：{ledger_row_selects} / 全部 SQL：{sqls}"
        assert len(ledger_selects) <= 2, (
            f"账本查询条数随行数增长了（N+1 又回来了）：{ledger_selects} / 全部 SQL：{sqls}"
        )
        lowered = ledger_row_selects[0].lower()
        assert "ledgers.entry_date >=" in lowered, (
            f"日期条件没有下推到 SQL（还是在 Python 里过滤）：{ledger_row_selects[0]}"
        )
        assert "ledgers.entry_date <=" in lowered, ledger_row_selects[0]

        ws = _sheet(content, "客户经营")
        got = _data_rows(ws)
        assert got == _buckets_as_pre_fix(db_session, s, e), (
            "客户经营导出的逐行结果与修复前不一致（日期过滤改坏了口径）"
        )

        # 我造的那几条：只有落在区间内的两笔算进来（边界含），区间外的两笔与临时户不算
        mine = [row for row in got if row[1] == (users["shipper"].full_name or users["shipper"].phone)]
        assert len(mine) == 1, f"货主桶不是一条：{mine}"
        assert mine[0][2] >= 2 and mine[0][3] >= 555.0, f"区间内的两笔没算进来：{mine}"
        joined = "\n".join(str(c) for row in got for c in row)
        assert f"{tag}区间外临时户" not in joined, "区间外的临时货主出现在这份报表里"
    finally:
        _purge(db_session, ledger_ids=[r.id for r in probes])


# ---------------------------------------------------------------- ⑤ R2-4 文件名与内容同源
def test_turnover_and_products_export_filenames_follow_the_real_window(client, token_dispatcher):
    """`turnover`/`products` 的文件名必须写**真实取数区间**（内容就是 `mode` 的窗口）。

    缺陷现场：这两个 kind 一律用锚点日命名，而 `mode=month` 的内容是整月 ——
    `?mode=month&date=2026-09-18` 导出的文件叫 `turnover-report-2026-09-18.xlsx`、
    里面却是 09-01~09-30。对账/存档时按文件名找回来会拿到一份"名字与内容不符"的凭证。
    """
    h = auth_headers(token_dispatcher)

    fn, content = _export(client, h, kind="turnover", mode="month", date="2026-09-18")
    assert fn == "turnover-report-2026-09-01_2026-09-30.xlsx", f"文件名没跟真实区间走：{fn}"
    # 名字与内容同源：正文的逐日序列确实覆盖到 09-30（内容是整月）
    labels = [row[0] for row in _data_rows(_sheet(content, "营业纵览"), min_row=9) if isinstance(row[0], str)]
    assert "9-30" in labels, f"正文没覆盖到月末（说明名字与内容还是两个口径）：{labels[:5]}…{labels[-3:]}"

    fn2, _ = _export(client, h, kind="products", mode="month", date="2026-09-18")
    assert fn2 == "products-report-2026-09-01_2026-09-30.xlsx", f"文件名没跟真实区间走：{fn2}"

    fn3, _ = _export(client, h, kind="turnover", mode="week", date="2026-09-18")
    assert fn3 == "turnover-report-2026-09-14_2026-09-20.xlsx", f"周窗口的文件名不对：{fn3}"

    fn4, _ = _export(client, h, kind="turnover", mode="day", date="2026-09-18")
    assert fn4 == "turnover-report-2026-09-18.xlsx", f"单日的文件名不对：{fn4}"

    # ⚠️ 不许回归：按区间导出的四个 kind 仍以 date_from/date_to 为准（第十五轮修的那条）
    fn5, _ = _export(
        client, h, kind="customers", mode="day", date="2026-09-18",
        date_from="2026-09-01", date_to="2026-09-05",
    )
    assert fn5 == "customers-report-2026-09-01_2026-09-05.xlsx", f"按区间导出的文件名错了：{fn5}"

    # ⚠️ 2026-09-22 更新（报表时间控件换成"档位药丸"那一轮）：**六个 kind 现在都认区间**。
    #    以前 `date_from/date_to` 对 turnover/products 不起作用，所以上面那条"名字跟真实区间走"
    #    只能拿 mode 的窗口来比；现在页面发的就是区间（`ReportFinance.windowOf`），
    #    区间**优先于 mode+anchor** —— 于是这里的正确期望变成：**名字 = 传进去的那一段**，
    #    内容也跟着那一段（不再是"名字写区间、内容按 mode"）。
    fn6, content6 = _export(
        client, h, kind="turnover", mode="month", date="2026-09-18",
        date_from="2026-09-01", date_to="2026-09-03",
    )
    assert fn6 == "turnover-report-2026-09-01_2026-09-03.xlsx", (
        f"turnover 的文件名没跟（现在生效的）区间走：{fn6}"
    )
    labels6 = [row[0] for row in _data_rows(_sheet(content6, "营业纵览"), min_row=9) if isinstance(row[0], str)]
    assert "9-4" not in labels6 and "9-3" in labels6, (
        f"内容没按这段区间取数（三天窗口里出现了区间外的日子）：{labels6}"
    )
    # 区间优先也体现在表头：写的是那一段，而不是「month 2026-09月」
    assert _sheet(content6, "营业纵览").cell(1, 2).value == "9-1~9-3", (
        _sheet(content6, "营业纵览").cell(1, 2).value
    )
