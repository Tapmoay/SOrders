"""经营利润表（FEAT-0011）：把已经算得出来的四块钱汇成一张「这月赚了多少」。

## 为什么要有这个文件（2026-10-04）

每一块钱本来都有唯一实现，但**从来没有一处把它们相减**：营业纵览只到「商品毛利」
（`cost_covered_amount - cost_total`），司机应得在运费结算页、开销在收支页 ——
老板问「这月赚了多少」，今天只能在三个页面之间手工相减。

`GET /reports/profit` 把四块钱按**同一个窗口**汇合，且前几格**直接取营业纵览同一批聚合数**：
于是「利润表的营业额 == 营业纵览的营业额」是构造上成立的，不是两处各算一遍再碰运气
（毛利上已经栽过一次：同一批数两种算法，同月 72,177.75 vs 正确 10,789.00）。

这个文件钉六件事：
1. 与营业纵览**逐分一致**（营业额 / 配送成本 / 商品成本 / 覆盖率行数）；
2. 恒等式当场成立（见 `app/schemas/reports.py::ProfitReportOut` 的注释）；
3. 期间费用**按 exp_date 落窗口**，分类明细之和 == 合计；
4. 税今天没有数据源 -> 恒为 0，且口径说明里写清「为什么是 0」（⛔ 不许编一个税率）；
5. 边界：半截窗口 / 反了的窗口 -> 400；空窗口 -> 全 0 且不炸；
6. 权限：不是派单员 -> 403（报表是全店口径，没有「只看自己那份」的版本）。
"""

from __future__ import annotations

import itertools
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from tests.conftest import auth_headers

PROBE = "利润表探针"
_SEQ = itertools.count(1)


def _get(client: TestClient, h: dict[str, str], path: str, **params) -> dict:
    q = "&".join(f"{k}={v}" for k, v in params.items() if v is not None)
    r = client.get(f"/api/v1/reports/{path}?{q}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _dec(d: dict, key: str) -> Decimal:
    return Decimal(str(d[key]))


def _mk_delivered(db_session, users, *, unit_price: str = "50.00", qty: int = 2,
                  cost: str = "30.00", freight: str = "20.00") -> tuple[int, str]:
    """造一张**今天的**已送达单（带成本快照 -> 毛利那一格非空），返回 (订单 id, 当地日)。

    ⚠️ 订单号带自增序号：同一个进程的测试库不会因为同一天重复建单而撞唯一键。
    """
    from app.core.business_time import business_today, utc_now_naive
    from app.models import Order, OrderProduct
    from app.models.enums import OrderStatus

    day = business_today()
    stamp = day.strftime("%m%d")
    order = Order(
        order_no=f"PROFIT{stamp}-{next(_SEQ)}",
        status=OrderStatus.DELIVERED,
        order_date=day,
        shipper_id=users["shipper"].id,
        driver_id=users["driver"].id,
        delivered_at=utc_now_naive(),
        dispatched_at=utc_now_naive(),
        freight_fee=Decimal(freight),
        driver_billing_mode_snapshot="PIECE",
        delivery_description=PROBE,
    )
    db_session.add(order)
    db_session.flush()
    db_session.add(
        OrderProduct(
            order_id=order.id,
            product_name_snapshot=PROBE,
            quantity=qty,
            unit_price=Decimal(unit_price),
            line_total=Decimal(unit_price) * qty,
            cost_price_snapshot=Decimal(cost),
        )
    )
    db_session.commit()
    return order.id, day.isoformat()


def _purge(db_session, order_id: int) -> None:
    from app.models import Order, OrderProduct

    db_session.query(OrderProduct).filter(OrderProduct.order_id == order_id).delete()
    db_session.query(Order).filter(Order.id == order_id).delete()
    db_session.commit()


# ------------------------------------------------------ ① 与营业纵览逐分一致


@pytest.mark.dispatcher
@pytest.mark.fast
@pytest.mark.regression
def test_利润表与营业纵览逐分一致(client, token_dispatcher, db_session, users):
    """利润表的收入/配送成本/商品成本必须**与营业纵览是同一批数**（差一分钱就是又算了一遍）。"""
    h = auth_headers(token_dispatcher)
    oid, day = _mk_delivered(db_session, users)
    try:
        p = _get(client, h, "profit", mode="day", date=day)
        t = _get(client, h, "turnover", mode="day", date=day)
        assert _dec(p, "revenue_total") == _dec(t, "total_amount"), "营业额与营业纵览对不上"
        assert _dec(p, "delivery_cost") == _dec(t, "total_freight"), "配送成本与营业纵览的司机运费支出对不上"
        assert _dec(p, "cost_total") == _dec(t, "cost_total"), "商品成本与营业纵览对不上"
        assert p["total_lines"] == t["total_lines"]
        assert p["covered_lines"] == t["cost_covered_lines"]
        assert p["cost_avg_lines"] == t["cost_avg_lines"]
        assert p["cost_snapshot_lines"] == t["cost_snapshot_lines"]
        assert _dec(p, "gross_profit") == _dec(t, "cost_covered_amount") - _dec(t, "cost_total"), (
            "毛利两侧不是同一批行（历史事故：同月 72,177.75 vs 正确 10,789.00）"
        )
        # 非空转：这一格确实有我们自己造的那张单（100 元收入 / 60 元成本 / 20 元司机应得）
        assert _dec(p, "revenue_total") >= Decimal("100.00")
        assert _dec(p, "cost_total") >= Decimal("60.00")
        assert _dec(p, "delivery_cost") >= Decimal("20.00")
    finally:
        _purge(db_session, oid)


# ------------------------------------------------------ ② 恒等式


@pytest.mark.dispatcher
@pytest.mark.fast
@pytest.mark.regression
def test_恒等式当场成立(client, token_dispatcher, db_session, users):
    """不参与毛利的收入要单列、毛利两侧同一批行、营业利润是四级相减、明细之和 == 合计。"""
    h = auth_headers(token_dispatcher)
    oid, day = _mk_delivered(db_session, users)
    try:
        p = _get(client, h, "profit", mode="day", date=day)
        rev, cov, unc = (_dec(p, k) for k in ("revenue_total", "revenue_covered", "revenue_uncovered"))
        assert rev == cov + unc, "不参与毛利的收入既没有被单列、也没进毛利（它会凭空消失）"
        gross, cost = _dec(p, "gross_profit"), _dec(p, "cost_total")
        assert gross == cov - cost, "毛利不是「参与毛利的收入 - 商品成本」"
        assert _dec(p, "operating_profit") == (
            gross
            - _dec(p, "delivery_cost")
            - _dec(p, "operating_expense_total")
            - _dec(p, "tax_total")
        ), "营业利润不是四级相减的结果"
        rows = p["operating_expenses"]
        assert sum(Decimal(str(r["amount"])) for r in rows) == _dec(p, "operating_expense_total"), (
            "分类明细之和与合计对不上（页面上两处会各说各话）"
        )
        amounts = [Decimal(str(r["amount"])) for r in rows]
        assert amounts == sorted(amounts, reverse=True), "明细没有按金额从大到小（第一眼要看到最大那笔）"
    finally:
        _purge(db_session, oid)


# ------------------------------------------------------ ③ 期间费用按发生日落窗口


@pytest.mark.dispatcher
@pytest.mark.fast
def test_期间费用按发生日落窗口(client, token_dispatcher):
    """开销按 `exp_date` 落窗口：今天那张只出现在今天这一段，2000 年那一段必须是 0 且没有明细。"""
    h = auth_headers(token_dispatcher)
    from app.core.business_time import business_today

    today = business_today().isoformat()
    r = client.post(
        "/api/v1/expenses",
        json={"exp_date": today, "category": "fuel", "amount": "12.34", "note": PROBE},
        headers=h,
    )
    assert r.status_code == 200, r.text
    inside = _get(client, h, "profit", mode="day", date=today)
    assert _dec(inside, "operating_expense_total") >= Decimal("12.34")
    hit = [x for x in inside["operating_expenses"] if Decimal(str(x["amount"])) >= Decimal("12.34")]
    assert hit, inside["operating_expenses"]
    outside = _get(client, h, "profit", mode="day", date="2000-01-15", date_from="2000-01-01", date_to="2000-01-31")
    assert _dec(outside, "operating_expense_total") == 0
    assert outside["operating_expenses"] == []


# ------------------------------------------------------ ④ 税与口径说明


@pytest.mark.dispatcher
@pytest.mark.fast
def test_税恒为0且口径说明写清了为什么(client, token_dispatcher):
    """税今天没有数据源（ACCOUNTING_V2 的 P3 未落地）-> 如实写 0，并把原因写在页面上。"""
    h = auth_headers(token_dispatcher)
    p = _get(client, h, "profit", mode="day", date="2026-10-04")
    assert _dec(p, "tax_total") == 0, "没有税账就不许编一个税率：这一格只能是 0"
    notes = p["notes"]
    assert notes, "口径说明不许是空的（老板会照这些数判断挣没挣钱）"
    assert any("税" in n for n in notes), notes
    assert any("折旧" in n for n in notes), notes
    assert any("工资" in n for n in notes), notes
    assert any("毛利" in n for n in notes), notes


# ------------------------------------------------------ ⑤ 边界：窗口


@pytest.mark.dispatcher
@pytest.mark.fast
def test_半截窗口与反了的窗口都拒绝(client, token_dispatcher):
    """只给一头 -> 400（不许猜另一头）；顺序反了 -> 400（不许安静地返回空集）。"""
    h = auth_headers(token_dispatcher)
    for params in ({"date_from": "2026-09-01"}, {"date_to": "2026-09-30"}):
        q = "&".join(f"{k}={v}" for k, v in params.items())
        r = client.get(f"/api/v1/reports/profit?mode=day&date=2026-09-22&{q}", headers=h)
        assert r.status_code == 400, r.text
    r = client.get(
        "/api/v1/reports/profit?mode=day&date=2026-09-22&date_from=2026-09-30&date_to=2026-09-01",
        headers=h,
    )
    assert r.status_code == 400, r.text


@pytest.mark.dispatcher
@pytest.mark.fast
def test_空窗口全0且不炸(client, token_dispatcher):
    """一段没有数据的窗口：四块钱全是 0、标签跟着窗口走、口径说明仍然在（空态也要说得清）。"""
    h = auth_headers(token_dispatcher)
    p = _get(client, h, "profit", mode="day", date="2000-01-15", date_from="2000-01-01", date_to="2000-01-31")
    assert p["period_label"] == "2000-01月"
    for key in (
        "revenue_total", "revenue_covered", "revenue_uncovered", "cost_total", "gross_profit",
        "delivery_cost", "operating_expense_total", "tax_total", "operating_profit",
        "collected", "arrears_total",
    ):
        assert _dec(p, key) == 0, key
    assert p["operating_expenses"] == []
    assert p["notes"], "空窗口也要给出说法"


# ------------------------------------------------------ ⑥ 权限与导出


@pytest.mark.dispatcher
@pytest.mark.fast
def test_不是派单员拿不到(client, token_shipper, token_driver):
    """报表是全店口径（营业额/毛利/司机成本），没有「只看自己那份」的版本 -> 403。"""
    for tok in (token_shipper, token_driver):
        r = client.get("/api/v1/reports/profit?mode=day&date=2026-09-22", headers=auth_headers(tok))
        assert r.status_code == 403, r.text


@pytest.mark.dispatcher
def test_导出kind_profit有经营利润那张表(client, token_dispatcher, db_session, users):
    """导出的第 7 个 kind：文件名写真实区间，表里四块钱 + 覆盖率 + 口径说明一个不少。"""
    h = auth_headers(token_dispatcher)
    oid, day = _mk_delivered(db_session, users)
    try:
        r = client.get(
            f"/api/v1/reports/export?kind=profit&mode=day&date={day}&date_from={day}&date_to={day}",
            headers=h,
        )
        assert r.status_code == 200, r.text
        disp = r.headers.get("content-disposition") or ""
        assert day in disp, disp
        ws = load_workbook(BytesIO(r.content))["经营利润"]
        cells = [str(ws.cell(i, j).value or "") for i in range(1, ws.max_row + 1) for j in range(1, 8)]
        for label in ("营业收入(应收)", "商品成本", "商品毛利", "配送成本(司机应得)",
                      "期间费用", "税金及附加", "营业利润", "成本覆盖率", "口径说明"):
            assert label in cells, (label, cells)
        assert any("入库加权平均进货价" in c for c in cells), "覆盖率那一段必须说清成本是从哪来的"
    finally:
        _purge(db_session, oid)

