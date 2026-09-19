"""第十三轮审计的回归测试（司机端全链 + 报表导出两轮红队审查的成果）。

| 编号 | 症状 | 这条测试钉什么 |
|---|---|---|
| R13-D1 | 司机能给**已进回收站**的在途单送达：读侧 404、写侧 200，库存照扣、账单照出，而这张单谁都看不见 | 已删除的单：送达/上传凭证/加备注/派单**全部**拒绝 |
| R13-D2 | 账单月份按 UTC、结算页按业务当地月 → 月初 8 小时的单落在错月份；AI 补单「已完成」却一张没生成 | 账单月份 = 业务当地月 |
| R13-D3 | 月份只校验格式 → 9 月就能造出 10 月的可支付工资单（本机库里真有一张 6500 元的 `2026-10`） | 未来月份一律拒绝 |
| R13-R3 | 撤销数按 UTC 日、日报小时桶按 UTC 小时 | 两者都按业务当地时区 |
| R13-R4 | 「挂账未收」两处判据不同（多一条 `payment_method == "arrears"`） | 只有 `paid=False` 一条判据 |
| R13-R7 | 导出金额写成文本（Excel 求和不到钱）、单均价 28 位小数 | 金额是**数字**、两位小数 |
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

from tests.conftest import auth_headers


def _mk_driver(client, h, name: str = "十三轮探针司机") -> tuple[int, str]:
    import uuid

    phone = "13" + uuid.uuid4().hex[:9].translate(str.maketrans("abcdef", "012345"))
    r = client.post(
        "/api/v1/users",
        json={"phone": phone, "password": "pass12345", "full_name": name, "role": "driver"},
        headers=h,
    )
    assert r.status_code in (200, 201), r.text
    did = int(r.json()["id"])
    r = client.post("/api/v1/auth/login", json={"phone": phone, "password": "pass12345"})
    assert r.status_code == 200, r.text
    return did, r.json()["access_token"]


def _deliverable_order(client, h, db_session, driver_id: int, driver_token: str, **over) -> dict:
    """建一张「已接单」的单（可送达状态）。"""
    payload = {
        "shipper_id": 2,
        "lines": [{"product_name_snapshot": "探针货", "quantity": 1, "unit_price": "100", "line_total": "100"}],
        "address_detail": "十三轮探针路 1 号",
        "freight_fee": "100",
    }
    payload.update(over)
    r = client.post("/api/v1/orders", json=payload, headers=h)
    assert r.status_code in (200, 201), r.text
    oid = r.json()["id"]
    assert client.post(f"/api/v1/orders/{oid}/assign", json={"driver_id": driver_id}, headers=h).status_code in (200, 201)
    assert client.post(f"/api/v1/orders/{oid}/driver-ack", headers=auth_headers(driver_token)).status_code in (200, 201)
    return {"id": oid}


def _soft_delete(db_session, order_id: int) -> None:
    from app.models import Order

    db_session.expire_all()
    o = db_session.get(Order, order_id)
    o.deleted_at = datetime(2026, 9, 19, 0, 0)
    db_session.commit()


# --------------------------------------------------- R13-D1 隔离区的单不许被司机写
def test_driver_cannot_complete_a_deleted_order(client, token_dispatcher, db_session):
    """读侧 404 的单，写侧也必须拒绝——否则会落成一张"谁都看不见的已送达单"。"""
    from app.models import Order, OrderStatus

    h = auth_headers(token_dispatcher)
    did, dtoken = _mk_driver(client, h)
    hd = auth_headers(dtoken)

    # ⚠️ 先证明"**没被删的**同一种单能正常送达"——否则这条测试可能只是被别的守卫挡住
    #    （第一版就踩到了：司机没挂计费规则 → 卡在"请至少上传一张送达照片" → 断言恒真）。
    ok_order = _deliverable_order(client, h, db_session, did, dtoken)
    body = {"payment": "arrears", "delivery_photo_urls": ["/static/uploads/delivery/probe.jpg"]}
    r = client.post(f"/api/v1/orders/{ok_order['id']}/complete", json=body, headers=hd)
    assert r.status_code in (200, 201), f"前提不成立：正常的单都送不达 → {r.status_code} {r.text}"

    o = _deliverable_order(client, h, db_session, did, dtoken)
    _soft_delete(db_session, o["id"])

    # 读侧：司机看到的是"订单不存在"（前提，否则这条测试没有对象）
    assert client.get(f"/api/v1/orders/{o['id']}", headers=hd).status_code == 404

    r = client.post(f"/api/v1/orders/{o['id']}/complete", json=body, headers=hd)
    assert r.status_code in (400, 404), f"已删除的单不该能送达：{r.status_code} {r.text}"
    db_session.expire_all()
    assert db_session.get(Order, o["id"]).status == OrderStatus.ACCEPTED, "被拒绝的送达却改了状态"

    # 同一族的另一个写口
    r = client.post(f"/api/v1/orders/{o['id']}/driver-note", json={"note": "试试"}, headers=hd)
    assert r.status_code in (400, 404, 403), r.text

    # 派单：不校验的话会给司机推一条"新派单"，而司机点进去是 404
    db_session.expire_all()
    row = db_session.get(Order, o["id"])
    row.status = OrderStatus.PENDING_DISPATCH
    db_session.commit()
    r = client.post(f"/api/v1/orders/{o['id']}/assign", json={"driver_id": did}, headers=h)
    assert r.status_code == 400, f"已删除的单不该能派单：{r.status_code} {r.text}"
    assert "删" in r.json()["detail"], r.json()


# ------------------------------------------- R13-D2 账单月份 = 业务当地月
def test_bill_month_uses_business_local_month():
    """UTC 月末 18:00（= 当地次月 02:00）的单，账单必须落在**当地那个月**。"""
    from datetime import timezone

    from app.services.accounting_service import _month_of

    # 当地 2026-10-01 02:00 == UTC 2026-09-30 18:00
    assert _month_of(datetime(2026, 9, 30, 18, 0, tzinfo=timezone.utc)) == "2026-10"
    # 库里读回来是 naive（SQLite），同样按 UTC 解释
    assert _month_of(datetime(2026, 9, 30, 18, 0)) == "2026-10"
    # 当地 2026-09-30 23:00 == UTC 15:00，还在 9 月
    assert _month_of(datetime(2026, 9, 30, 15, 0, tzinfo=timezone.utc)) == "2026-09"


# ------------------------------------------- R13-D3 未来月份不许生成账单
def test_future_month_rejected_for_bills_and_settlements(client, token_dispatcher):
    from app.core.business_time import business_local
    from datetime import timezone

    h = auth_headers(token_dispatcher)
    now_month = business_local(datetime.now(timezone.utc)).strftime("%Y-%m")
    y, m = int(now_month[:4]), int(now_month[5:])
    future = f"{y + 1 if m == 12 else y:04d}-{1 if m == 12 else m + 1:02d}"

    r = client.post("/api/v1/driver-bills/generate", json={"month": future}, headers=h)
    assert r.status_code in (400, 422), f"未来月份不该能生成工资单：{r.status_code} {r.text}"
    assert "不能晚于本月" in r.text, r.text

    # 本月仍然可以（别把正常用法一起挡住）
    r = client.post("/api/v1/driver-bills/generate", json={"month": now_month}, headers=h)
    assert r.status_code in (200, 201), r.text


# ------------------------------------------- R13-R3 撤销数与小时桶按当地时区
def test_turnover_cancelled_count_uses_business_day(client, token_dispatcher, db_session):
    """当地 09-19 凌晨撤销的单要算进 09-19（原来按 UTC 日算进 09-18，那天的数字整行消失）。"""
    from app.api.v1.reports import build_turnover
    from app.models import Order, OrderStatus

    h = auth_headers(token_dispatcher)
    d19 = date(2026, 9, 19)
    before = build_turnover(db_session, "day", d19)["cancelled_orders"]

    r = client.post(
        "/api/v1/orders",
        json={
            "shipper_id": 2,
            "lines": [{"product_name_snapshot": "撤销探针", "quantity": 1, "unit_price": "10", "line_total": "10"}],
            "address_detail": "十三轮撤销探针路",
        },
        headers=h,
    )
    oid = r.json()["id"]
    db_session.expire_all()
    o = db_session.get(Order, oid)
    o.status = OrderStatus.CANCELLED
    # UTC 2026-09-18 20:00 == 当地 2026-09-19 04:00
    o.cancelled_at = datetime(2026, 9, 18, 20, 0)
    db_session.commit()

    after = build_turnover(db_session, "day", d19)["cancelled_orders"]
    assert after == before + 1, f"当地 09-19 撤销的单没算进 09-19（{before} → {after}）"


def test_turnover_hour_bucket_uses_business_hour(client, token_dispatcher, db_session):
    """当地 09-19 04:00 送达的单要落进「04时」桶，而不是 UTC 的「20时」。"""
    from app.api.v1.reports import build_turnover
    from app.models import Order, OrderStatus

    h = auth_headers(token_dispatcher)
    r = client.post(
        "/api/v1/orders",
        json={
            "shipper_id": 2,
            "lines": [{"product_name_snapshot": "小时探针", "quantity": 1, "unit_price": "10", "line_total": "10"}],
            "address_detail": "十三轮小时探针路",
        },
        headers=h,
    )
    oid = r.json()["id"]
    db_session.expire_all()
    o = db_session.get(Order, oid)
    o.status = OrderStatus.DELIVERED
    o.delivered_at = datetime(2026, 9, 18, 20, 0)  # 当地 09-19 04:00
    db_session.commit()

    series = build_turnover(db_session, "day", date(2026, 9, 19))["series"]
    labels = {s.label: s.orders for s in series}
    assert labels.get("04时", 0) >= 1, f"当地 4 点送达的单没进「04时」：{labels}"
    assert labels.get("20时", 0) == 0, f"按 UTC 分桶了（出现「20时」）：{labels}"


# ------------------------------------------- R13-R4 挂账未收只有一条判据
def test_arrears_definitions_agree(client, token_dispatcher, db_session):
    """营业纵览的「挂账未收」与客户经营页/导出的那一份必须是同一个数。

    构造一张 `cash + paid=0 + collect_cash=0` 的历史形态单（老数据/导库都可能长这样）——
    两处判据不同时它会只被算进其中一边（实测差 ¥85.50 / 7 张）。
    """
    from app.api.v1.reports import build_arrears_summary, build_turnover
    from app.models import Order, OrderProduct, OrderStatus, Product

    h = auth_headers(token_dispatcher)
    prod = Product(name="挂账口径探针", default_unit_price=Decimal("50"), stock=10)
    db_session.add(prod)
    db_session.flush()
    o = Order(
        order_no="SOARREARS1301",
        shipper_id=2,
        status=OrderStatus.DELIVERED,
        order_date=date(2026, 9, 19),
        delivered_at=datetime(2026, 9, 19, 6, 0),
        address_detail="十三轮挂账探针路",
        freight_fee=Decimal("0"),
        # 关键：现金收款方式 + 未收款 —— 新代码不会再写这种组合，但库里可能已经有
        payment_method="cash",
        collect_cash=False,
        paid=False,
        arrears_unit_name="十三轮挂账单位",
    )
    db_session.add(o)
    db_session.flush()
    db_session.add(
        OrderProduct(
            order_id=o.id, product_id=prod.id, product_name_snapshot=prod.name,
            quantity=1, unit_price=Decimal("50"), line_total=Decimal("50"),
        )
    )
    db_session.commit()

    turn = build_turnover(db_session, "day", date(2026, 9, 19))
    summary = build_arrears_summary(db_session, date(2026, 9, 19), date(2026, 9, 19))
    unit = next((x for x in summary if x["name"] == "十三轮挂账单位"), None)
    assert unit is not None, f"客户经营那一份没算进这张单：{summary}"
    assert Decimal(str(unit["amount"])) == Decimal("50.00")
    # 营业纵览那一份把这张单算进"挂账未收"（paid=False 一条判据）
    assert turn["arrears_total"] >= Decimal("50.00")


# ------------------------------------------- R13-R7 导出金额是数字
def test_export_amounts_are_numbers(tmp_path):
    """导出的金额必须是**数字**（Excel 里能 SUM），且单均价不超过两位小数。"""
    from io import BytesIO

    from openpyxl import Workbook

    # 直接测写入函数的语义：`_money` 把 Decimal 变成两位小数的 float
    from app.api.v1.reports import _money

    assert _money(Decimal("97131.7500")) == 97131.75
    assert isinstance(_money("193.4895418326693227091633466"), float)
    assert _money("193.4895418326693227091633466") == 193.49
    assert _money(None) == ""

    # 端到端：把 `_money` 的结果塞进 openpyxl，确认它是数值型而不是字符串
    wb = Workbook()
    ws = wb.active
    ws.append([_money(Decimal("100.00"))])
    buf = BytesIO()
    wb.save(buf)
    assert ws["A1"].value == 100.0
    assert isinstance(ws["A1"].value, (int, float))
