"""报表的时间窗口：**给了区间就按区间，没给就按 mode+anchor**——两条路必须算出同一份数。

## 为什么要有这个文件（2026-09-22）

App 的报表时间控件换成了**档位药丸**（今天 / 昨天 / 近 7 天 / 本月 / 上月 / 自定义），
它发给后端的是一段**明确区间**（`date_from` + `date_to`）。而报表原来只有
`mode`（day/week/month）+ `anchor` 一条路 —— 于是：

* 六个页签里**营业纵览 / 商品经营**走 mode+anchor、其余四个走 date_range，
  **同一屏两个窗口口径**（这本来就是老毛病：标题说一段、数字是另一段）；
* 「近 7 天」「自定义」这两种窗口**根本表达不出来**（没有对应的 mode）——
  也就是说"给报表换个时间控件"这件事，绕不开后端补一个区间入口。

补法必须是**共用同一段聚合**（`build_turnover` / `build_products` 只多收一个 `span`），
否则就成了"两张表两套算法"，而页面上完全看不出来。这个文件就是钉这一点：

1. **同一个窗口、两条路 → 逐项相等**（含 `series` 每一天、毛利/成本/货损/资金那几列）；
2. **区间优先于 mode+anchor**（两个都给时按区间走）；
3. **只给一头 → 400**（半截窗口要么查全量要么查出空列表，两种都不是用户要的）；
4. `period_label` 与窗口一致（整月 `2026-09月` / 一段 `9-1~9-20` / 一天 `9-18`）；
5. 导出（`kind=turnover|products`）也认区间 —— 文件名的日期段 = **真实取数的那一段**。
"""

from __future__ import annotations

from datetime import timedelta
from decimal import Decimal

import pytest
from starlette.testclient import TestClient

from tests.conftest import auth_headers

PROBE = "报表窗口探针"


def _get(client: TestClient, h: dict[str, str], path: str, **params) -> dict:
    q = "&".join(f"{k}={v}" for k, v in params.items() if v is not None)
    r = client.get(f"/api/v1/reports/{path}?{q}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _money_keys(d: dict) -> dict:
    """所有金额列（JSON 里是字符串）→ Decimal；其余数字/列表原样比。"""
    out = {}
    for k, v in d.items():
        if k == "_window" or k == "period_label":
            continue
        if isinstance(v, str):
            try:
                out[k] = Decimal(v)
            except Exception:  # noqa: BLE001 - 认不出是金额的字符串（比如日期）就不比
                continue
        elif isinstance(v, (int, float, list)):
            out[k] = v
    return out


def _mk_delivered(client: TestClient, db_session, users) -> tuple[int, str]:
    """造一张**今天的**已送达单（报表窗口按业务当地日算），返回 (订单 id, 当地日)。"""
    from app.core.business_time import business_today, utc_now_naive
    from app.models import Order, OrderProduct
    from app.models.enums import OrderStatus

    day = business_today()
    order = Order(
        order_no=f"RPTWIN{day.strftime('%m%d')}",
        status=OrderStatus.DELIVERED,
        order_date=day,
        shipper_id=users["shipper"].id,
        driver_id=users["driver"].id,
        delivered_at=utc_now_naive(),
        dispatched_at=utc_now_naive(),
        freight_fee=Decimal("20.00"),
        driver_billing_mode_snapshot="PIECE",
        delivery_description=PROBE,
    )
    db_session.add(order)
    db_session.flush()
    db_session.add(
        OrderProduct(
            order_id=order.id,
            product_name_snapshot=PROBE,
            quantity=2,
            unit_price=Decimal("50.00"),
            line_total=Decimal("100.00"),
        )
    )
    db_session.commit()
    return order.id, day.isoformat()


def _purge(db_session, order_id: int) -> None:
    from app.models import Order, OrderProduct

    db_session.query(OrderProduct).filter(OrderProduct.order_id == order_id).delete()
    db_session.query(Order).filter(Order.id == order_id).delete()
    db_session.commit()


# ---------------------------------------------------------------- ① 两条路同一个窗口


@pytest.mark.dispatcher
def test_month_window_equals_explicit_range(client, token_dispatcher, db_session, users):
    """`mode=month&date=D` 与「本月那一段区间」必须**逐项相等**（含 series 的每一天）。"""
    h = auth_headers(token_dispatcher)
    oid, day = _mk_delivered(client, db_session, users)
    try:
        d = day  # 2026-09-22 这种
        first = d[:8] + "01"
        # 本月最后一天：交给后端自己算（`mode=month` 的那条路），区间那边用同一段
        by_mode = _get(client, h, "turnover", mode="month", date=d)
        label = by_mode["period_label"]  # 形如 2026-09月
        assert label.endswith("月"), label
        # 该月最后一天 = 下月 1 日往前一天
        from datetime import date

        y, m = int(d[:4]), int(d[5:7])
        nxt = date(y + (1 if m == 12 else 0), 1 if m == 12 else m + 1, 1)
        last = (nxt - timedelta(days=1)).isoformat()
        by_range = _get(client, h, "turnover", mode="day", date=d, date_from=first, date_to=last)
        assert _money_keys(by_mode) == _money_keys(by_range), (
            "同一个窗口、两条路算出了两份数（页面会跟着请求用什么参数而变）"
        )
        # 非空转：这个月里确实有我们自己造的那张单
        assert Decimal(by_mode["total_amount"]) >= Decimal("100.00")
    finally:
        _purge(db_session, oid)


@pytest.mark.dispatcher
def test_day_window_equals_explicit_range_and_products_too(client, token_dispatcher, db_session, users):
    """按日和「当天那一段」相等；**商品经营**那一份也必须相等（两页同一个窗口口径）。"""
    h = auth_headers(token_dispatcher)
    oid, day = _mk_delivered(client, db_session, users)
    try:
        for path in ("turnover", "products"):
            by_mode = _get(client, h, path, mode="day", date=day)
            by_range = _get(client, h, path, mode="week", date=day, date_from=day, date_to=day)
            assert _money_keys(by_mode) == _money_keys(by_range), path
            assert Decimal(by_mode["total_amount"]) >= Decimal("100.00"), path
    finally:
        _purge(db_session, oid)


@pytest.mark.dispatcher
def test_range_overrides_mode_and_anchor(client, token_dispatcher, db_session, users):
    """两个都给时**按区间走**（否则页面上选了"近 7 天"、后端还按 mode 算一个月）。"""
    h = auth_headers(token_dispatcher)
    oid, day = _mk_delivered(client, db_session, users)
    try:
        # mode=month 会含这一天；区间故意取一个**不含它**的窗口（2000 年）→ 必须是 0
        both = _get(client, h, "turnover", mode="month", date=day, date_from="2000-01-01", date_to="2000-01-31")
        assert Decimal(both["total_amount"]) == 0, "区间没有优先于 mode+anchor"
        assert both["period_label"] == "2000-01月"
    finally:
        _purge(db_session, oid)


# ---------------------------------------------------------------- ② 半截窗口 / 反了的窗口


@pytest.mark.dispatcher
def test_only_one_end_given_is_rejected(client, token_dispatcher):
    """只给一头 → 400（不许猜另一头：半截窗口要么查全量要么查出空列表）。"""
    h = auth_headers(token_dispatcher)
    for params in ({"date_from": "2026-09-01"}, {"date_to": "2026-09-30"}):
        r = client.get("/api/v1/reports/turnover?mode=day&date=2026-09-22&" + "&".join(f"{k}={v}" for k, v in params.items()), headers=h)
        assert r.status_code == 400, r.text


@pytest.mark.dispatcher
def test_reversed_range_is_rejected(client, token_dispatcher):
    h = auth_headers(token_dispatcher)
    r = client.get(
        "/api/v1/reports/turnover?mode=day&date=2026-09-22&date_from=2026-09-30&date_to=2026-09-01",
        headers=h,
    )
    assert r.status_code == 400, r.text


# ---------------------------------------------------------------- ③ 标签与窗口一致


@pytest.mark.dispatcher
def test_period_label_matches_the_window(client, token_dispatcher):
    """整月写 `2026-09月`、一段写 `9-1~9-20`、一天写 `9-18` —— 界面上那几个数必须对得上这段。"""
    h = auth_headers(token_dispatcher)
    whole = _get(client, h, "turnover", mode="day", date="2026-09-18", date_from="2026-09-01", date_to="2026-09-30")
    assert whole["period_label"] == "2026-09月"
    part = _get(client, h, "turnover", mode="day", date="2026-09-18", date_from="2026-09-01", date_to="2026-09-20")
    assert part["period_label"] == "9-1~9-20"
    one = _get(client, h, "turnover", mode="day", date="2026-09-18", date_from="2026-09-18", date_to="2026-09-18")
    assert one["period_label"] == "9-18"


# ---------------------------------------------------------------- ④ 导出也认区间


@pytest.mark.dispatcher
def test_turnover_export_accepts_range_and_names_file_by_it(client, token_dispatcher):
    """`kind=turnover` 用区间导出时：文件名与表头都写**真实取数的那一段**，不是锚点日。

    以前 `date_from/date_to` 对 turnover/products 是**被忽略**的（只有四个 kind 认）——
    于是"报表页按区间看、导出按 mode 导"就会导出另一段时间的数据，还提示"成功"。
    """
    h = auth_headers(token_dispatcher)
    r = client.get(
        "/api/v1/reports/export?kind=turnover&mode=day&date=2026-09-18"
        "&date_from=2026-09-01&date_to=2026-09-20",
        headers=h,
    )
    assert r.status_code == 200, r.text
    disp = r.headers.get("content-disposition") or ""
    assert "2026-09-01_2026-09-20" in disp, disp
    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(r.content))["营业纵览"]
    assert ws.cell(1, 2).value == "9-1~9-20", ws.cell(1, 2).value
